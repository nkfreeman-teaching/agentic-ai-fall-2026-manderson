"""Render the shared segmentation results as figures, a DOCX report, and HTML explorer."""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import seaborn as sns
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent
OUT = ROOT / "output" / "full"
RESULT = json.loads((OUT / "results.json").read_text(encoding="utf-8"))
SEGMENTS = RESULT["segments"]
SOURCE = RESULT["source"]
MODEL = RESULT["model"]
COLORS = ["#167d8d", "#ce7130", "#6d5ca6", "#7b8c29"]
NAMES = ["Broader baskets", "Value-oriented baskets"]
ORAL = {row["segment"]: row for row in RESULT["offer_evidence"]["oral_hygiene"]}
MEAT = {row["segment"]: row for row in RESULT["offer_evidence"]["private_meat"]}
CAMPAIGNS = {row["campaign_type"]: row for row in RESULT["campaign_context"]}
OFFERS = [
    "Test an oral-hygiene offer against a same-value general nonfuel offer. "
    f"Oral-hygiene purchases appear in {ORAL[1]['buying_households']:,} of {SEGMENTS[0]['households']:,} households ({ORAL[1]['buying_households'] / SEGMENTS[0]['households']:.1%}) in this group, "
    f"compared with {ORAL[2]['buying_households']:,} of {SEGMENTS[1]['households']:,} ({ORAL[2]['buying_households'] / SEGMENTS[1]['households']:.1%}) in the other group. "
    f"Buyers in this group made a median of {ORAL[1]['median_baskets_among_buyers']:.0f} oral-hygiene purchase trips. "
    "This is an affinity signal, not evidence that a coupon will add sales.",
    "Test a private-label meat or packaged-meat offer against a same-value general nonfuel offer. "
    f"Private-label meat purchases appear in {MEAT[2]['buying_households']:,} of {SEGMENTS[1]['households']:,} households ({MEAT[2]['buying_households'] / SEGMENTS[1]['households']:.1%}) in this group, "
    f"compared with {MEAT[1]['buying_households']:,} of {SEGMENTS[0]['households']:,} ({MEAT[1]['buying_households'] / SEGMENTS[0]['households']:.1%}) in the other group. "
    f"Buyers in this group made a median of {MEAT[2]['median_baskets_among_buyers']:.0f} such purchase trips. "
    "The pattern does not establish incremental response.",
]


def figure_setup() -> None:
    sns.set_style("whitegrid")
    plt.rcParams.update({
        "font.family": "DejaVu Serif",
        "font.size": 14,
        "axes.labelsize": 14,
        "xtick.labelsize": 12,
        "ytick.labelsize": 12,
        "legend.fontsize": 11,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "savefig.dpi": 180,
    })


def save_figures() -> None:
    figure_setup()
    candidates = MODEL["candidates"]
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.plot(
        [row["k"] for row in candidates],
        [row["silhouette"] for row in candidates],
        marker="o",
        color=COLORS[0],
    )
    ax.set(xlabel="Number of segments", ylabel="Mean silhouette score", xticks=list(range(2, 9)))
    ax.set_ylim(0, max(row["silhouette"] for row in candidates) * 1.18)
    fig.tight_layout()
    fig.savefig(OUT / "cluster_selection.png")
    plt.close(fig)

    labels = ["Private label", "Retail discounts", "Drug and general merchandise", "Meat", "Packaged meat"]
    keys = ["private_label_share", "retail_discount_share", "share_drug_gm", "share_meat", "share_meat_pckgd"]
    y = np.arange(len(labels))
    fig, ax = plt.subplots(figsize=(8, 4))
    width = 0.36
    for index, segment in enumerate(SEGMENTS):
        values = [100 * segment[key] for key in keys]
        ax.barh(
            y + (index - 0.5) * width,
            values,
            height=width,
            color=COLORS[index],
            edgecolor="black",
            linewidth=0.4,
            label=NAMES[index],
        )
    ax.set(yticks=y, yticklabels=labels, xlabel="Median household share (%)")
    ax.invert_yaxis()
    ax.legend(frameon=False, loc="lower right")
    fig.tight_layout()
    fig.savefig(OUT / "segment_profile.png")
    plt.close(fig)


def add_reference(paragraph, label: str, url: str) -> None:
    paragraph.add_run(f"{label}. {url}")


def make_report() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name in ["Title", "Heading 1", "Heading 2"]:
        styles[name].font.name = "Aptos Display"
        styles[name].font.color.rgb = RGBColor(25, 56, 70)
    doc.add_heading("Customer segmentation proof of concept", 0)
    p = doc.add_paragraph("Complete Journey simulated retail data | 22 September 2026")
    p.style = "Subtitle"
    doc.add_heading("Decision summary", 1)
    doc.add_paragraph(
        f"The files contain {SOURCE['purchase_lines']:,} purchase lines, {SOURCE['baskets']:,} baskets, and "
        f"{SOURCE['households']:,} transacting households over roughly one year. "
        f"We could form nonfuel shopping profiles for {SOURCE['model_households']:,} households. "
        f"The remaining {SOURCE['fuel_only_households']} purchased fuel but had no non-fuel purchases. "
        "A two-group solution separates basket composition and price-related behavior more than visit frequency. "
        "They motivate oral-hygiene and private-label meat coupon tests, without evidence of incremental coupon sales."
    )
    doc.add_paragraph(
        "The publisher identifies this Python release as simulated educational data. Its household demographics cover "
        f"{SOURCE['demographic_households']:,} of the {SOURCE['model_households']:,} profiled households "
        f"({SOURCE['demographic_coverage']:.1%}). Thus, neither the group proportions nor the proposed offers describe a real retailer's customers."
    )
    doc.add_heading("What the groups show", 1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Shading Accent 1"
    for cell, label in zip(table.rows[0].cells, ["Group", "Households", "Nonfuel baskets", "Basket receipts", "Private label", "Retail discount"]):
        cell.text = label
    for index, segment in enumerate(SEGMENTS):
        values = [
            NAMES[index],
            f"{segment['households']:,} ({segment['share']:.1%})",
            f"{segment['basket_count']:.0f}",
            f"${segment['mean_basket_receipts']:.2f}",
            f"{segment['private_label_share']:.1%}",
            f"{segment['retail_discount_share']:.1%}",
        ]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    doc.add_paragraph("Values are household medians except group size. Receipts are retailer receipts, not a direct measure of customer outlay.")
    doc.add_picture(str(OUT / "segment_profile.png"), width=Inches(5.2))
    doc.add_paragraph(
        "Figure 1. Median shares across households. The value-oriented group allocates more non-fuel receipts to private-label goods and records a larger retail-discount share. "
        "Each percentage has its own denominator, and median category shares need not sum to 100%."
    )
    doc.add_heading("How the groups were formed", 1)
    doc.add_paragraph(
        "We linked purchase lines to product metadata and aggregated them to households. "
        "The inputs capture nonfuel basket frequency, mean basket receipts, days since the last purchase, "
        "private-label share, retail-discount share, and spending shares across major non-fuel departments. "
        "We excluded fuel from the grouping because fuel-only trips describe a different purchase occasion. "
        "Demographics and coupon redemptions did not determine membership."
    )
    doc.add_paragraph(
        "Recency, frequency, and monetary measures have a history in customer-base analysis (Fader, Hardie, and Lee 2005). "
        "We use related descriptive features here, but do not estimate customer lifetime value."
    )
    doc.add_paragraph(
        "We applied a log transform to skewed activity measures, clipped each feature at the first and 99th percentiles, "
        "scaled by its interquartile range, and gave activity, value orientation, and product mix equal total weight. "
        "K-means models with two through eight groups were compared using silhouette separation, repeated 80% household subsamples, "
        "and a minimum group size of 5%. The highest eligible silhouette selected two groups."
    )
    chosen = next(row for row in MODEL["candidates"] if row["k"] == MODEL["selected_k"])
    doc.add_paragraph(
        f"The selected mean silhouette is {chosen['silhouette']:.3f}, which indicates modest separation. "
        f"Mean adjusted Rand agreement across subsamples is {chosen['stability_ari']:.3f}. "
        f"Adding fuel share changes {MODEL['fuel_inclusion_sensitivity']['changed_households']} of {SOURCE['model_households']:,} assignments after matching labels. "
        f"Removing households with fewer than four nonfuel baskets yields adjusted Rand agreement of {MODEL['sparse_history_sensitivity']['agreement_ari']:.3f} on the retained households. "
        "These diagnostics support a simple descriptive split but do not establish a natural or permanent customer typology."
    )
    doc.add_picture(str(OUT / "cluster_selection.png"), width=Inches(4.8))
    doc.add_paragraph("Figure 2. Mean silhouette score by number of groups. The first point marks the selected two-group solution.")
    doc.add_heading("Range in household activity", 2)
    quantiles = SOURCE["profile_quantiles"]
    spread = doc.add_table(rows=1, cols=4)
    spread.style = "Light Shading Accent 1"
    for cell, label in zip(spread.rows[0].cells, ["Measure", "10th percentile", "Median", "90th percentile"]):
        cell.text = label
    for values in [
        ["Nonfuel baskets", f"{quantiles['baskets_p10']:.0f}", f"{SOURCE['median_baskets']:.0f}", f"{quantiles['baskets_p90']:.0f}"],
        ["Nonfuel receipts", f"${quantiles['receipts_p10']:,.0f}", f"${SOURCE['median_receipts']:,.0f}", f"${quantiles['receipts_p90']:,.0f}"],
    ]:
        for cell, value in zip(spread.add_row().cells, values):
            cell.text = value
    doc.add_paragraph("The wide range in purchase activity motivates robust scaling and limits confidence for households with short histories.")
    doc.add_page_break()
    doc.add_heading("Coupon tests", 1)
    for index, offer in enumerate(OFFERS):
        doc.add_heading(NAMES[index], 2)
        doc.add_paragraph(offer)
    doc.add_paragraph(
        "Freeze group assignments using a pre-test purchase period, then randomize households concurrently within each group to the category offer, "
        "a general offer with comparable eligibility and cost, or no offer for eight weeks. Calculate the sample size before launch. "
        "Record offer delivery, eligibility, redemption, all shopping, discount cost, and true gross margin. "
        "The primary comparison is incremental gross margin per assigned household for the category offer versus no offer. "
        "The general offer is a secondary contrast. Gopalakrishnan and Park (2021) found coupon effects can arise without redemption, "
        "which is why the test must measure all shopping. Daljord et al. (2023) show why optional redemption complicates promotion evaluation."
    )
    doc.add_paragraph(
        "Campaign-household assignments and coupon lists permit descriptive reach calculations for some historical campaigns. "
        f"The files contain {CAMPAIGNS['Type B']['household_campaign_assignments']:,} Type B and "
        f"{CAMPAIGNS['Type C']['household_campaign_assignments']:,} Type C household-campaign assignments. "
        f"All {SOURCE['coupon_redemptions']:,} redemption records match a campaign assignment and fall within campaign dates. "
        "Type A's household-specific coupon selection is not fully recorded, and none of these observational records identifies incremental offer effects."
    )
    doc.add_heading("Data limits", 1)
    doc.add_paragraph(
        f"No matching product record exists for {SOURCE['unmatched_product_lines']:,} purchase lines, "
        f"and {SOURCE['households_under_four_baskets']:,} profiled households have fewer than four non-fuel baskets. "
        "The first are retained in an other category. The second have limited individual histories, so their assignments are less secure. "
        "Historical campaign records permit descriptive checks but lack random assignment and gross margin for a causal response estimate. "
        "The simulated source also prevents external validation with real retail behavior."
    )
    doc.add_heading("Sources", 1)
    refs = [
        ("Complete Journey Python documentation, data notice", "https://cunningjames.github.io/completejourney_py/user-guide/datasets/"),
        ("Fader, Hardie, and Lee (2005), RFM and CLV", "https://doi.org/10.1509/jmkr.2005.42.4.415"),
        ("Gopalakrishnan and Park (2021), The Impact of Coupons on the Visit-to-Purchase Funnel", "https://doi.org/10.1287/mksc.2020.1232"),
        ("Daljord et al. (2023), The Design and Targeting of Compliance Promotions", "https://doi.org/10.1287/mksc.2022.1420"),
    ]
    for label, url in refs:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.style = "Normal"
        add_reference(paragraph, label, url)
        for run in paragraph.runs:
            run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.text = "Educational proof of concept | Simulated data"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(OUT / "customer_segmentation_report.docx")


def make_explorer() -> None:
    data = json.dumps({
        "segments": SEGMENTS,
        "source": SOURCE,
        "model": MODEL,
        "campaign_context": RESULT["campaign_context"],
        "names": NAMES,
        "offers": OFFERS,
    }).replace("</", "<\\/")
    html = r'''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Customer segments | Complete Journey</title>
<style>
:root{--ink:#173040;--muted:#5a6e77;--paper:#f6f8f7;--teal:#167d8d;--orange:#ce7130;--line:#d9e2e3}
*{box-sizing:border-box}body{margin:0;background:var(--paper);color:var(--ink);font:16px/1.55 system-ui,Arial,sans-serif}
header{background:#123a48;color:white;padding:35px max(22px,calc((100vw - 1120px)/2));border-bottom:7px solid #57bdc1}
.eyebrow{text-transform:uppercase;letter-spacing:.14em;font-size:12px;font-weight:700;color:#a9d8d9}h1{font-size:clamp(30px,4vw,48px);line-height:1.08;margin:8px 0 12px}header p{max-width:760px;color:#d6e6e8;margin:0}
main{max-width:1120px;margin:auto;padding:26px 22px 65px}.notice{background:#e8f2ef;border-left:4px solid var(--teal);padding:14px 18px;margin-bottom:24px}
.stats{display:grid;grid-template-columns:repeat(4,1fr);gap:12px}.stat,.card{background:white;border:1px solid var(--line);border-radius:11px;box-shadow:0 5px 18px #15323d0b}.stat{padding:16px}.stat strong{display:block;font-size:28px;line-height:1.2}.stat span{font-size:13px;color:var(--muted)}
section{margin-top:32px}h2{font-size:25px;line-height:1.2;margin:0 0 12px}h3{font-size:18px;margin:0 0 9px}.sub{color:var(--muted);margin-top:0;max-width:820px}.grid{display:grid;grid-template-columns:1fr 1fr;gap:18px}.card{padding:22px}
.controls{display:flex;gap:12px;flex-wrap:wrap;margin:16px 0}label{font-size:14px;font-weight:700}.controls label{width:320px;max-width:100%;min-width:0}select{display:block;width:100%;min-width:0;font:inherit;padding:9px 32px 9px 10px;border:1px solid #9fb8bd;border-radius:7px;background:white}
.bars{display:grid;gap:18px;margin-top:14px}.bar-row{display:grid;grid-template-columns:205px 1fr 78px;align-items:center;gap:11px;font-size:14px}.track{height:22px;background:#e8eeee;border-radius:4px;overflow:hidden}.fill{height:100%;border-radius:4px}.bar-row b{text-align:right;font-variant-numeric:tabular-nums}
.segment-switch{display:flex;gap:8px;flex-wrap:wrap}.segment-switch button{padding:10px 14px;border:1px solid var(--line);background:white;color:var(--ink);border-radius:7px;cursor:pointer;font:inherit}.segment-switch button.active{background:var(--ink);color:white}.segment-switch button:focus-visible,select:focus-visible{outline:3px solid #f2a64e;outline-offset:2px}
.metric-list{display:grid;grid-template-columns:repeat(3,1fr);gap:10px}.metric{border-top:1px solid var(--line);padding-top:10px}.metric b{font-size:21px;display:block}.metric small{color:var(--muted)}
table{border-collapse:collapse;width:100%;font-size:14px}td,th{padding:9px 10px;text-align:left;border-bottom:1px solid var(--line)}th{background:#eaf0f0}td:nth-child(n+2),th:nth-child(n+2){text-align:right;font-variant-numeric:tabular-nums}
.foot{font-size:13px;color:var(--muted)}a{color:#087086}footer{margin-top:38px;padding-top:20px;border-top:1px solid var(--line);font-size:13px;color:var(--muted)}
@media(max-width:760px){.stats{grid-template-columns:repeat(2,1fr)}.grid{grid-template-columns:1fr}.bar-row{grid-template-columns:125px 1fr 55px}.metric-list{grid-template-columns:repeat(2,1fr)}.card{padding:17px}table{font-size:12px}}
@media(max-width:500px){.campaign-table thead{display:none}.campaign-table,.campaign-table tbody,.campaign-table tr,.campaign-table td{display:block;width:100%}.campaign-table tr{padding:10px 0;border-bottom:1px solid var(--line)}.campaign-table td{display:flex;justify-content:space-between;gap:10px;padding:3px 0;border:0;text-align:right}.campaign-table td:before{content:attr(data-label);font-weight:700;text-align:left}.campaign-table td:first-child{font-weight:700}}
</style></head><body>
<header><div class="eyebrow">Educational retail analytics</div><h1>Customer segments</h1><p>A working example of how retail purchase histories can guide the design of a coupon test. The underlying Complete Journey release is simulated.</p></header>
<main><div class="notice"><strong>Interpretation boundary.</strong> These profiles describe patterns in simulated data. Coupon ideas below are testable hypotheses, not estimates of sales lift.</div>
<div class="stats" id="stats"></div>
<section><h2>How many groups?</h2><p class="sub">Two groups give the strongest eligible silhouette score. Their separation is modest, so the labels describe tendencies rather than fixed customer types.</p><div class="grid"><div class="card"><h3>Group sizes</h3><div id="size-bars" class="bars"></div></div><div class="card"><h3>Selection check</h3><div id="selection"></div><p class="foot">Silhouette rewards separation. The stability score is adjusted Rand agreement across repeated 80% household subsamples.</p></div></div></section>
<section><h2>Compare shopping patterns</h2><p class="sub">Choose a measure to compare household medians. Retailer receipts include the accounting effects of coupons and are not direct customer outlay.</p><div class="controls"><label>Measure<select id="metric-select"><option value="private_label_share">Private-label share</option><option value="retail_discount_share">Retail discount share</option><option value="share_drug_gm">Drug and general merchandise share</option><option value="share_meat_combined">Meat and packaged meat share</option><option value="mean_basket_receipts">Mean basket receipts</option><option value="basket_count">Nonfuel baskets</option></select></label></div><div class="card"><div id="comparison" class="bars" aria-live="polite"></div></div></section>
<section><h2>Explore a group</h2><div id="segment-switch" class="segment-switch"></div><div class="grid" style="margin-top:14px"><div class="card"><h3 id="segment-name"></h3><div id="segment-metrics" class="metric-list"></div></div><div class="card"><h3>Offer to test</h3><p id="offer"></p><p class="foot">Freeze groups before the test. Randomize households within each group to the proposed offer, a comparable general offer, or no offer. Compare gross margin per assigned household over eight weeks. Calculate sample size before launch.</p></div></div></section>
<section><h2>Historical campaign context</h2><p class="sub">Campaign assignments describe reach, not incremental sales from an offer.</p><div class="card"><table class="campaign-table"><thead><tr><th>Campaign type</th><th>Campaigns</th><th>Household-campaign assignments</th><th>Distinct assigned households</th></tr></thead><tbody id="campaign-table"></tbody></table><p class="foot" id="redemption-note"></p></div></section>
<section><h2>Method and limits</h2><div class="grid"><div class="card"><h3>Method</h3><p>Purchase lines were joined to product metadata and summarized by household. Nonfuel visit frequency, basket size, recency, private-label share, retail-discount share, and department mix shaped the groups. Skewed activity measures were transformed, extreme values clipped, and three feature domains given equal weight. K-means models with two to eight groups were compared.</p></div><div class="card"><h3>What these data cannot show</h3><p>Demographics describe only a subset of households and did not form the groups. Type B and C campaign assignments permit descriptive reach checks, but historical offers were not randomized. The files do not provide gross margin. These records cannot establish offer effects or support commercial decisions.</p></div></div><p class="foot" id="data-limits"></p></section>
<footer>Sources: <a href="https://cunningjames.github.io/completejourney_py/user-guide/datasets/">Complete Journey documentation</a>; <a href="https://doi.org/10.1509/jmkr.2005.42.4.415">Fader et al. (2005)</a>; <a href="https://doi.org/10.1287/mksc.2020.1232">Gopalakrishnan and Park (2021)</a>; <a href="https://doi.org/10.1287/mksc.2022.1420">Daljord et al. (2023)</a>. Figures are computed from the local parquet files.</footer></main>
<script>const D=__DATA__;const fmt=new Intl.NumberFormat('en-US');const pct=x=>(x*100).toFixed(1)+'%';const money=x=>'$'+x.toFixed(2);const colors=['#167d8d','#ce7130'];
document.getElementById('stats').innerHTML=[['Purchase lines',fmt.format(D.source.purchase_lines)],['Transacting households',fmt.format(D.source.households)],['Nonfuel profiles',fmt.format(D.source.model_households)],['Demographic coverage',pct(D.source.demographic_coverage)]].map(([l,v])=>`<div class="stat"><strong>${v}</strong><span>${l}</span></div>`).join('');
function bars(container,values,labels,format){const max=Math.max(...values)*1.08;document.getElementById(container).innerHTML=values.map((v,i)=>`<div class="bar-row"><span>${labels[i]}</span><div class="track"><div class="fill" style="width:${100*v/max}%;background:${colors[i]}"></div></div><b>${format(v)}</b></div>`).join('')}
bars('size-bars',D.segments.map(s=>s.households),D.names,v=>fmt.format(v));
const chosen=D.model.candidates.find(c=>c.k===D.model.selected_k);document.getElementById('selection').innerHTML=`<div class="metric-list"><div class="metric"><b>${D.model.selected_k}</b><small>Selected groups</small></div><div class="metric"><b>${chosen.silhouette.toFixed(3)}</b><small>Mean silhouette</small></div><div class="metric"><b>${chosen.stability_ari.toFixed(3)}</b><small>Stability</small></div></div><p class="foot">The smallest group holds ${pct(chosen.smallest_share)} of profiled households.</p>`;
const measures={private_label_share:pct,retail_discount_share:pct,share_drug_gm:pct,share_meat_combined:pct,mean_basket_receipts:money,basket_count:v=>v.toFixed(0)};function updateComparison(){const key=document.getElementById('metric-select').value;bars('comparison',D.segments.map(s=>s[key]),D.names,measures[key])}document.getElementById('metric-select').addEventListener('change',updateComparison);updateComparison();
function showSegment(i){document.querySelectorAll('#segment-switch button').forEach((b,j)=>{b.classList.toggle('active',j===i);b.setAttribute('aria-pressed',j===i)});const s=D.segments[i];document.getElementById('segment-name').textContent=D.names[i];document.getElementById('offer').textContent=D.offers[i];document.getElementById('segment-metrics').innerHTML=[['Households',fmt.format(s.households)],['Nonfuel baskets',s.basket_count.toFixed(0)],['Basket receipts',money(s.mean_basket_receipts)],['Private label',pct(s.private_label_share)],['Retail discount',pct(s.retail_discount_share)],['Meat and packaged meat',pct(s.share_meat_combined)]].map(([l,v])=>`<div class="metric"><b>${v}</b><small>${l}</small></div>`).join('')}
document.getElementById('segment-switch').innerHTML=D.names.map((name,i)=>`<button type="button" data-i="${i}">${name}</button>`).join('');document.querySelectorAll('#segment-switch button').forEach(b=>b.addEventListener('click',()=>showSegment(Number(b.dataset.i))));showSegment(0);
document.getElementById('campaign-table').innerHTML=D.campaign_context.map(c=>`<tr><td data-label="Campaign type">${c.campaign_type}</td><td data-label="Campaigns">${fmt.format(c.campaigns)}</td><td data-label="Household-campaign assignments">${fmt.format(c.household_campaign_assignments)}</td><td data-label="Distinct assigned households">${fmt.format(c.assigned_households)}</td></tr>`).join('');
document.getElementById('redemption-note').textContent=`Redemption events are separate: ${D.campaign_context.map(c=>`${c.campaign_type}, ${fmt.format(c.redemption_events)}`).join('; ')}. These counts are not response rates. For Type B and C, the publisher says assigned households received the campaign coupons. Type A coupon selection was personalized and is not fully recorded at household level. ${fmt.format(D.source.redemptions_without_campaign_assignment)} redemptions lack an assignment match and ${fmt.format(D.source.redemptions_outside_campaign_dates)} fall outside campaign dates.`;
document.getElementById('data-limits').textContent=`${fmt.format(D.source.unmatched_product_lines)} purchase lines have no matching product record. ${fmt.format(D.source.households_under_four_baskets)} profiled households have fewer than four nonfuel baskets. ${fmt.format(D.source.fuel_only_households)} fuel-only households have no nonfuel profile. Removing sparse shoppers gave ${D.model.sparse_history_sensitivity.agreement_ari.toFixed(3)} adjusted Rand agreement on retained households.`;
</script></body></html>'''
    (OUT / "segment_explorer.html").write_text(html.replace("__DATA__", data), encoding="utf-8")


if __name__ == "__main__":
    save_figures()
    make_report()
    make_explorer()
    print("Rendered report and explorer in", OUT)
